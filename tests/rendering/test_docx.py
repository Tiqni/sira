from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm, RGBColor

from sira.models.agents.output import Education, Project
from sira.rendering.docx import write_docx
from sira.rendering.errors import RenderError
from sira.rendering.templates import CLASSIC, COMPACT, MODERN, TemplateSpec
from tests.factories import make_cv


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#"))


@pytest.mark.parametrize("spec", [MODERN, CLASSIC, COMPACT])
def test_docx_styles_follow_the_spec(tmp_path: Path, spec: TemplateSpec):
    out = tmp_path / "r.docx"
    write_docx(make_cv(), spec, out)
    doc = Document(str(out))

    assert doc.styles["Normal"].font.name == spec.docx_font
    assert doc.styles["Heading 2"].font.color.rgb == _rgb(spec.heading_hex)
    assert doc.styles["Title"].font.color.rgb == _rgb(spec.accent_hex)
    assert bool(doc.styles["Heading 2"].font.all_caps) == (spec.heading_style == "caps")
    section = doc.sections[0]
    assert abs(section.left_margin - Mm(spec.margin_mm)) < Mm(0.1)
    assert abs(section.page_width - Mm(210)) < Mm(0.1)
    assert len(doc.tables) == 0  # ATS guard: single column, no tables

    for name in ("Title", "Heading 2"):
        rfonts = doc.styles[name].element.rPr.find(qn("w:rFonts"))
        assert not any(
            k.endswith("Theme") or k.endswith("theme") for k in rfonts.attrib
        )
        assert rfonts.get(qn("w:ascii")) == spec.docx_font

    # A section heading must never end a page on its own (issue #21).
    assert doc.styles["Heading 2"].paragraph_format.keep_with_next is True


def test_docx_document_order_and_text(tmp_path: Path):
    cv = make_cv().model_copy(
        update={
            "projects": [
                Project(name="Tool", description="A **CLI**", link="https://t.io")
            ],
            "education": [
                Education(
                    degree="BSc", institution="TU", dates="2016", details="Honours"
                )
            ],
        }
    )
    out = tmp_path / "r.docx"
    write_docx(cv, MODERN, out)
    doc = Document(str(out))

    paragraphs = [(p.style.name, p.text) for p in doc.paragraphs]
    assert paragraphs[0] == ("Title", "Jane Doe")
    assert paragraphs[1] == ("Contact", "jane@example.com")
    headings = [text for style, text in paragraphs if style == "Heading 2"]
    assert headings == ["Summary", "Skills", "Experience", "Projects", "Education"]
    assert ("Normal", "Languages: Python, SQL") in paragraphs
    assert ("Normal", "Engineer — Acme · 2022-2026") in paragraphs
    assert ("List Bullet", "Built services") in paragraphs
    assert ("Normal", "Tool — A CLI (https://t.io)") in paragraphs
    assert ("Normal", "BSc — TU · 2016") in paragraphs
    assert ("Normal", "Honours") in paragraphs

    bold_runs = [r.text for p in doc.paragraphs for r in p.runs if r.bold]
    assert "CLI" in bold_runs and "Engineer" in bold_runs and "Languages: " in bold_runs
    hyperlinks = [
        r.target_ref for r in doc.part.rels.values() if "hyperlink" in r.reltype
    ]
    assert hyperlinks == ["https://t.io"]


def test_docx_bold_inside_link_text_becomes_a_bold_hyperlink_run(tmp_path: Path):
    cv = make_cv().model_copy(
        update={"certifications": ["[**CKAD:** Certified Kubernetes](https://c.io)"]}
    )
    out = tmp_path / "r.docx"
    write_docx(cv, MODERN, out)
    doc = Document(str(out))
    paragraph = next(p for p in doc.paragraphs if "Certified Kubernetes" in p.text)
    hyperlink_runs = paragraph._p.findall(".//" + qn("w:hyperlink") + "/" + qn("w:r"))
    texts = [r.find(qn("w:t")).text for r in hyperlink_runs]
    bold = [r.find(qn("w:rPr")).find(qn("w:b")) is not None for r in hyperlink_runs]
    assert texts == ["CKAD:", " Certified Kubernetes"]
    assert bold == [True, False]
    assert "**" not in paragraph.text


def test_docx_failure_raises_render_error(tmp_path: Path):
    blocker = tmp_path / "file"
    blocker.write_text("not a directory")
    with pytest.raises(RenderError, match="DOCX"):
        write_docx(make_cv(), MODERN, blocker / "r.docx")
