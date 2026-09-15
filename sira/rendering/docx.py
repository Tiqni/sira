"""CV → DOCX with python-docx, styled from a ``TemplateSpec``.

Only paragraphs and character runs are used — no tables or text boxes — so
applicant tracking systems read the file top to bottom.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from sira.models.agents.output import CV
from sira.rendering.errors import RenderError
from sira.rendering.inline import Run, inline_markdown_to_runs
from sira.rendering.templates import TemplateSpec

_HYPERLINK_REL = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
)
_CODE_FONT = "Consolas"


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _set_font(style, name: str, size_pt: float, hex_color: str, *, bold=None) -> None:
    style.font.name = name
    style.font.size = Pt(size_pt)
    style.font.color.rgb = _rgb(hex_color)
    if bold is not None:
        style.font.bold = bold
    # python-docx sets only w:ascii/w:hAnsi; fill every slot so Word never substitutes.
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)
    # A w:*Theme attribute overrides its named-font sibling per OOXML, and the
    # default template's Title/Heading 2 styles carry these — drop them so our
    # font actually wins instead of silently losing to the theme font.
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(attr), None)


def _remove_paragraph_borders(style) -> None:
    ppr = style.element.get_or_add_pPr()
    for border in ppr.findall(qn("w:pBdr")):
        ppr.remove(border)


def _set_bottom_border(style, hex_color: str) -> None:
    ppr = style.element.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")  # eighths of a point → 0.75pt, same as the CSS rule
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), hex_color.lstrip("#").upper())
    pbdr.append(bottom)
    ppr.append(pbdr)


def _set_letter_spacing(style, twentieths_of_pt: int) -> None:
    rpr = style.element.get_or_add_rPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), str(twentieths_of_pt))
    rpr.append(spacing)


def _apply_styles(doc, spec: TemplateSpec) -> None:
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Mm(spec.margin_mm))

    normal = doc.styles["Normal"]
    _set_font(normal, spec.docx_font, spec.base_pt, spec.text_hex)
    normal.paragraph_format.line_spacing = spec.line_height
    normal.paragraph_format.space_after = Pt(spec.item_gap_pt)

    title = doc.styles["Title"]
    _set_font(title, spec.docx_font, spec.name_pt, spec.accent_hex, bold=True)
    _remove_paragraph_borders(title)
    title.paragraph_format.space_after = Pt(2)
    # The default template's Title style carries a leftover letter-spacing
    # value that is not part of TemplateSpec — drop it so nothing but our
    # explicit heading_style="caps" spacing ever applies.
    rpr = title.element.rPr
    spacing = rpr.find(qn("w:spacing")) if rpr is not None else None
    if spacing is not None:
        rpr.remove(spacing)

    heading = doc.styles["Heading 2"]
    _set_font(heading, spec.docx_font, spec.heading_pt, spec.heading_hex, bold=True)
    _remove_paragraph_borders(heading)
    _set_bottom_border(heading, spec.heading_hex)
    heading.paragraph_format.space_before = Pt(spec.section_gap_pt)
    heading.paragraph_format.space_after = Pt(spec.item_gap_pt)
    heading.font.all_caps = spec.heading_style == "caps"
    if spec.heading_style == "caps":
        _set_letter_spacing(heading, 20)  # 1pt, same as the CSS letter-spacing

    doc.styles["List Bullet"].paragraph_format.space_after = Pt(1)

    contact = doc.styles.add_style("Contact", WD_STYLE_TYPE.PARAGRAPH)
    contact.base_style = normal
    _set_font(contact, spec.docx_font, spec.base_pt, spec.muted_hex)
    contact.paragraph_format.space_after = Pt(spec.section_gap_pt)

    dates = doc.styles.add_style("Dates", WD_STYLE_TYPE.CHARACTER)
    dates.font.color.rgb = _rgb(spec.muted_hex)


def _add_hyperlink(
    paragraph: Paragraph, run: Run, accent_hex: str, hyperlink=None
) -> "OxmlElement":
    """Append ``run`` as a hyperlink run; reuse ``hyperlink`` to extend one link.

    Returns the ``w:hyperlink`` element so consecutive runs of the same link
    ("<b>CKAD:</b> Certified") end up inside one hyperlink, not several.
    """
    if hyperlink is None:
        r_id = paragraph.part.relate_to(run.href, _HYPERLINK_REL, is_external=True)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), r_id)
        paragraph._p.append(hyperlink)
    docx_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    # OOXML wants rPr children in schema order: rFonts, b, i, …, color, u.
    if run.code:
        rfonts = OxmlElement("w:rFonts")
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), _CODE_FONT)
        rpr.append(rfonts)
    if run.bold:
        rpr.append(OxmlElement("w:b"))
    if run.italic:
        rpr.append(OxmlElement("w:i"))
    color = OxmlElement("w:color")
    color.set(qn("w:val"), accent_hex.lstrip("#").upper())
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    docx_run.append(rpr)
    text_el = OxmlElement("w:t")
    text_el.text = run.text
    text_el.set(qn("xml:space"), "preserve")
    docx_run.append(text_el)
    hyperlink.append(docx_run)
    return hyperlink


def _add_runs(paragraph: Paragraph, runs: list[Run], spec: TemplateSpec) -> None:
    hyperlink = None  # the open w:hyperlink element, shared by consecutive runs
    open_href = ""
    for run in runs:
        if run.href:
            if run.href != open_href:
                hyperlink, open_href = None, run.href
            hyperlink = _add_hyperlink(paragraph, run, spec.accent_hex, hyperlink)
            continue
        hyperlink, open_href = None, ""
        docx_run = paragraph.add_run(run.text)
        docx_run.bold = run.bold or None
        docx_run.italic = run.italic or None
        if run.code:
            docx_run.font.name = _CODE_FONT


def _add_inline(paragraph: Paragraph, text: str, spec: TemplateSpec) -> Paragraph:
    _add_runs(paragraph, inline_markdown_to_runs(text), spec)
    return paragraph


def _add_head_line(doc, *, lead: str, rest: str, dates: str) -> None:
    """``<b>lead</b> — rest · dates`` on one line (experience and education)."""
    paragraph = doc.add_paragraph()
    paragraph.add_run(lead).bold = True
    paragraph.add_run(f" — {rest}")
    if dates:
        paragraph.add_run(" · ")
        paragraph.add_run(dates, style="Dates")
    paragraph.paragraph_format.space_after = Pt(1)


def write_docx(cv: CV, spec: TemplateSpec, output_path: Path) -> None:
    """Write ``cv`` as a styled single-column DOCX.

    Raises:
        RenderError: on any python-docx or filesystem failure.
    """
    try:
        doc = Document()
        _apply_styles(doc, spec)

        doc.add_paragraph(cv.full_name, style="Title")
        items = cv.contact.display_items()
        if items:
            paragraph = doc.add_paragraph(style="Contact")
            for index, item in enumerate(items):
                if index:
                    paragraph.add_run(" · ")
                _add_inline(paragraph, item, spec)

        if cv.summary.strip():
            doc.add_paragraph("Summary", style="Heading 2")
            _add_inline(doc.add_paragraph(), cv.summary.strip(), spec)

        groups = [group for group in cv.skill_groups if group.skills]
        if groups:
            doc.add_paragraph("Skills", style="Heading 2")
            for group in groups:
                paragraph = doc.add_paragraph()
                paragraph.add_run(f"{group.category}: ").bold = True
                _add_inline(paragraph, ", ".join(group.skills), spec)
                paragraph.paragraph_format.space_after = Pt(1)

        if cv.experience:
            doc.add_paragraph("Experience", style="Heading 2")
            for exp in cv.experience:
                _add_head_line(doc, lead=exp.role, rest=exp.company, dates=exp.dates)
                for item in exp.highlights:
                    _add_inline(doc.add_paragraph(style="List Bullet"), item, spec)

        if cv.projects:
            doc.add_paragraph("Projects", style="Heading 2")
            for project in cv.projects:
                paragraph = doc.add_paragraph()
                paragraph.add_run(project.name).bold = True
                paragraph.add_run(" — ")
                _add_inline(paragraph, project.description, spec)
                if project.link:
                    paragraph.add_run(" (")
                    _add_inline(paragraph, project.link, spec)
                    paragraph.add_run(")")

        if cv.education:
            doc.add_paragraph("Education", style="Heading 2")
            for edu in cv.education:
                _add_head_line(
                    doc, lead=edu.degree, rest=edu.institution, dates=edu.dates
                )
                if edu.details:
                    paragraph = _add_inline(doc.add_paragraph(), edu.details, spec)
                    for run in paragraph.runs:
                        run.font.color.rgb = _rgb(spec.muted_hex)

        for heading, entries in (
            ("Certifications", cv.certifications),
            ("Publications", cv.publications),
        ):
            if entries:
                doc.add_paragraph(heading, style="Heading 2")
                for item in entries:
                    _add_inline(doc.add_paragraph(style="List Bullet"), item, spec)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
    except Exception as exc:  # noqa: BLE001 — wrap every library/IO failure
        raise RenderError(f"Failed to write DOCX: {exc}") from exc
